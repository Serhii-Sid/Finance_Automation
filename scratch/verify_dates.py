import openpyxl
import pandas as pd
import datetime

output_path = 'output/Total_Ledger.xlsx'

print("--- Reading Daily_Dashboard_Data via Pandas ---")
df_data = pd.read_excel(output_path, sheet_name='Daily_Dashboard_Data')
print(f"Total rows in Daily_Dashboard_Data: {len(df_data)}")

null_dates_count = df_data['Дата'].isna().sum()
detail_rows_count = df_data['Місяць'].astype(str).str.contains('↳').sum()
rahom_rows_count = df_data['Місяць'].astype(str).str.startswith('РАЗОМ').sum()

print(f"Null dates count: {null_dates_count}")
print(f"Detail sub-rows count: {detail_rows_count}")
print(f"RAHOM summary rows count: {rahom_rows_count}")

print("\n--- Reading via OpenPyXL ---")
wb = openpyxl.load_workbook(output_path, data_only=True)
sheet = wb['Daily_Dashboard_Data']

print(f"Sheet max_row: {sheet.max_row}")
date_col_idx = None
for col_idx in range(1, sheet.max_column + 1):
    header = sheet.cell(row=1, column=col_idx).value
    if header == 'Дата':
        date_col_idx = col_idx
        break

print(f"Column 'Дата' index: {date_col_idx}")
if date_col_idx:
    for row_idx in range(2, 6):
        cell = sheet.cell(row=row_idx, column=date_col_idx)
        val_col1 = sheet.cell(row=row_idx, column=1).value
        print(f"Row {row_idx}: Month='{val_col1}', Date value={cell.value!r}, type={cell.data_type!r}, format={cell.number_format!r}")
